# -*- coding: utf-8 -*-
"""造 docx / xlsx / pptx 测试样本（内容已知，用于精确验证提取结果）。
用法： python make-office-fixtures.py <输出目录>

注意：样本要放在 test-fixtures/ 而不是 test-artifacts/ ——
后者是冒烟测试的派生目录，smoke-driver.js 每次开跑都会整个删掉重建。
"""
import os
import sys
import shutil

out = sys.argv[1] if len(sys.argv) > 1 else 'test-fixtures/office'
os.makedirs(out, exist_ok=True)
made = []

# ---------- docx ----------
try:
    from docx import Document

    doc = Document()
    doc.add_heading('季度报告 Quarter Report', level=1)
    doc.add_paragraph('第一段：中文内容测试 123。')
    doc.add_paragraph('特殊字符 & < > " \' 以及全角（括号）')
    doc.add_paragraph('')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '表头A'
    table.cell(0, 1).text = 'Header B'
    table.cell(1, 0).text = '值1'
    table.cell(1, 1).text = 'value 2'
    doc.add_paragraph('最后一段结尾。')
    p = os.path.join(out, 'sample.docx')
    doc.save(p)
    made.append(p)
except Exception as e:
    print('docx 生成失败:', e)

# ---------- pptx ----------
try:
    from pptx import Presentation

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = '演示标题 Slide Title'
    if len(s1.placeholders) > 1:
        s1.placeholders[1].text = '副标题 bullet one'
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = '第二页'
    if len(s2.placeholders) > 1:
        s2.placeholders[1].text_frame.text = '要点一'
    p = os.path.join(out, 'sample.pptx')
    prs.save(p)
    made.append(p)
except Exception as e:
    print('pptx 生成失败:', e)

# ---------- xlsx ----------
try:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = '成绩'
    ws['A1'] = '姓名'
    ws['B1'] = '分数'
    ws['A2'] = '张三'
    ws['B2'] = 42
    ws['A3'] = 'Li Si'
    ws['B3'] = 99.5
    ws2 = wb.create_sheet('第二表')
    ws2['A1'] = '表二内容'
    p = os.path.join(out, 'sample.xlsx')
    wb.save(p)
    made.append(p)
except Exception as e:
    print('xlsx 生成失败(openpyxl 缺失?):', e)

# ---------- 真实第三方文件（python-docx / python-pptx 自带的空白模板） ----------
# 从已安装的包里定位模板，不写死路径
def package_template(module_name, rel):
    try:
        mod = __import__(module_name)
        base = os.path.dirname(os.path.abspath(mod.__file__))
        p = os.path.join(base, rel)
        return p if os.path.exists(p) else None
    except Exception as e:
        print('定位 %s 模板失败: %s' % (module_name, e))
        return None


for src, dst in [
    (package_template('docx', os.path.join('templates', 'default.docx')), 'real-template.docx'),
    (package_template('pptx', os.path.join('templates', 'default.pptx')), 'real-template.pptx'),
]:
    try:
        if src and os.path.exists(src):
            target = os.path.join(out, dst)
            shutil.copyfile(src, target)
            made.append(target)
    except Exception as e:
        print('复制失败', src, e)

# ---------- 老格式 .doc（只是改名的文本，验证不会崩） ----------
with open(os.path.join(out, 'legacy.doc'), 'wb') as f:
    f.write('这是一个假的旧版 doc 文件，不是 ZIP 结构。'.encode('utf-8'))

print('生成了 %d 个样本：' % len(made))
for m in made:
    print('  ', m, os.path.getsize(m), 'bytes')
